"""The Excel quality review: structure, no respondent rows, values equal the CSVs, formulas recalculate."""

import json
import shutil
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from nhanes_diabetes.config import Config, load_config
from nhanes_diabetes.workbook import (
    SHEETS,
    WORKBOOK,
    build_workbook,
    formula_values,
    verify_workbook,
)

REPO = Path(__file__).resolve().parents[1]
TABLES = REPO / "outputs" / "tables"
pytestmark = pytest.mark.skipif(not (TABLES / "undiagnosis_estimates.csv").exists(), reason="outputs not generated")

BANNER_START = "NHANES 2017-2018 public-use survey data"
FORBIDDEN_HEADERS = {"seqn", "respondent", "respondent_id", "weight", "wtmec2yr", "psu", "sdmvpsu", "stratum",
                     "sdmvstra", "age", "hba1c", "diq010", "bmi", "ridageyr"}


def csv(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / f"{name}.csv")


def table(wb, sheet: str, name: str) -> pd.DataFrame:
    """Read an Excel table back into a frame using only the sheet's own header row and cells."""
    ws = wb[sheet]
    first_col, first_row, last_col, last_row = range_boundaries(ws.tables[name].ref)
    header = [ws.cell(first_row, c).value for c in range(first_col, last_col + 1)]
    rows = [[ws.cell(r, c).value for c in range(first_col, last_col + 1)] for r in range(first_row + 1, last_row + 1)]
    return pd.DataFrame(rows, columns=header)


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> tuple[Config, Path]:
    config = load_config(root=REPO)
    return config, build_workbook(config, tmp_path_factory.mktemp("xl") / "review.xlsx")


@pytest.fixture(scope="module")
def wb(built):
    return load_workbook(built[1])


def test_the_seven_required_sheets_exist_in_order(wb) -> None:
    assert wb.sheetnames == ["Executive_Summary", "Cohort_Flow", "Care_Gaps", "Subgroup_Review", "Model_Tradeoffs",
                             "Data_Quality", "Metric_Dictionary"] == SHEETS


def test_every_sheet_has_the_banner_a_table_frozen_headers_and_flags(wb) -> None:
    for ws in wb.worksheets:
        assert str(ws["A1"].value).startswith(BANNER_START), ws.title
        assert "not a certified quality measure" in str(ws["A1"].value)
        assert ws.tables, f"{ws.title} has no Excel table"
        assert ws.freeze_panes, f"{ws.title} has no frozen panes"
        assert ws.conditional_formatting, f"{ws.title} has no warning styles"


def test_no_respondent_level_fields_or_row_counts(wb) -> None:
    for ws in wb.worksheets:
        for name in ws.tables:
            headers = {str(c).lower() for c in table(wb, ws.title, name).columns}
            assert not (headers & FORBIDDEN_HEADERS), f"{ws.title}.{name}: {headers & FORBIDDEN_HEADERS}"
        assert ws.max_row < 250, f"{ws.title} is too long to be a summary"


def test_care_gaps_reconcile_to_the_estimate_tables(wb) -> None:
    sheet = table(wb, "Care_Gaps", "CareGaps").set_index("group")
    est = csv("undiagnosis_estimates").set_index("group")
    summary = csv("group_summary").set_index("group_name")
    assert set(sheet.index) == set(est.index)
    for g in est.index:
        assert sheet.loc[g, "hba1c_positive_respondents"] == est.loc[g, "respondents"]
        assert sheet.loc[g, "estimate"] == pytest.approx(est.loc[g, "estimate"], abs=1e-12)
        assert sheet.loc[g, "ci_low"] == pytest.approx(est.loc[g, "ci_low_logit"], abs=1e-12)
        assert sheet.loc[g, "ci_high"] == pytest.approx(est.loc[g, "ci_high_logit"], abs=1e-12)
        assert sheet.loc[g, "effective_n_kish"] == pytest.approx(est.loc[g, "effective_n_kish"], abs=1e-9)
        assert sheet.loc[g, "weighted_undiagnosed"] == pytest.approx(summary.loc[g, "weighted_undiagnosed"], rel=1e-12)
        assert sheet.loc[g, "undiagnosed_respondents"] == summary.loc[g, "undiagnosed_among_positive"]
        assert (sheet.loc[g, "small_sample_warning"] == "SMALL SAMPLE") == bool(est.loc[g, "small_n_warning"])
    assert sheet.loc["Other/Multiracial", "small_sample_warning"] == "SMALL SAMPLE"


def test_subgroup_review_carries_holm_labels(wb, built) -> None:
    sheet = table(wb, "Subgroup_Review", "SubgroupReview").set_index("group")
    con = csv("undiagnosis_contrasts").set_index("group")
    for g in con.index:
        assert sheet.loc[g, "difference"] == pytest.approx(con.loc[g, "difference"], abs=1e-12)
        assert sheet.loc[g, "p_value_holm"] == pytest.approx(con.loc[g, "p_value_holm"], abs=1e-12)
        assert sheet.loc[g, "analysis_type"] == con.loc[g, "analysis_type"]
    values = formula_values(built[1])
    first = wb["Subgroup_Review"].tables["SubgroupReview"].ref.split(":")[0]
    row0 = int(first[1:]) + 1
    labels = {sheet.index[i]: values[("Subgroup_Review", f"Q{row0 + i}")] for i in range(len(sheet))}
    assert labels["Non-Hispanic White"] == "Reference group"
    assert labels["Non-Hispanic Black"] == "Pre-specified comparison"
    for g in ("Mexican American", "Other Hispanic", "Non-Hispanic Asian", "Other/Multiracial"):
        assert labels[g].startswith("Exploratory")
        assert ("below" if con.loc[g, "p_value_holm"] < 0.05 else "not below") in labels[g]


def test_model_tradeoffs_reconcile_to_the_model_tables(wb) -> None:
    comp = table(wb, "Model_Tradeoffs", "ModelComparison").set_index("model")
    auc, diff = csv("auc"), csv("auc_difference").set_index("model")
    for model in diff.index:
        crit = auc[(auc["model"] == model) & (auc["evaluated_against"] == "hba1c_criterion")].set_index("label")
        assert comp.loc[model, "auc_diagnosed_label"] == pytest.approx(crit.loc["diagnosed", "auc"], abs=1e-12)
        assert comp.loc[model, "auc_hba1c_label"] == pytest.approx(crit.loc["hba1c_pos", "auc"], abs=1e-12)
        assert comp.loc[model, "auc_difference"] == pytest.approx(diff.loc[model, "difference"], abs=1e-12)
        assert comp.loc[model, "p_value_bootstrap"] == pytest.approx(diff.loc[model, "p_value_bootstrap"], abs=1e-12)
    perf = table(wb, "Model_Tradeoffs", "SubgroupPerformance")
    sub = csv("subgroup_metrics")
    assert len(perf) == len(sub)
    merged = perf.merge(sub, on=["model", "group", "label"], suffixes=("_wb", ""))
    assert len(merged) == len(sub)
    for col in ("sensitivity", "sensitivity_ci_low", "specificity", "specificity_ci_high", "ppv", "npv"):
        assert (merged[f"{col}_wb"] - merged[col]).abs().max() < 1e-12, col
    burden = table(wb, "Model_Tradeoffs", "FalsePositiveBurden").merge(csv("specificity_cost"), on=["model", "group"], suffixes=("_wb", ""))
    assert len(burden) == len(csv("specificity_cost"))
    assert (burden["cost_wb"] - burden["cost"]).abs().max() < 1e-12


def test_cohort_flow_and_data_quality_match_their_sources(wb) -> None:
    flow = table(wb, "Cohort_Flow", "CohortFlow")
    att = csv("cohort_attrition")
    assert flow["remaining"].tolist() == att["n"].tolist()
    assert flow["removed"].tolist() == att["removed"].tolist()
    checks = table(wb, "Data_Quality", "QualityChecks")
    quality = csv("data_quality")
    assert checks["check"].tolist() == quality["check_name"].tolist()
    assert checks["status"].tolist() == ["PASS" if p else "FAIL" for p in quality["passed"]]
    miss = table(wb, "Data_Quality", "Missingness")
    src = csv("missingness_by_group")
    assert miss["missing"].tolist() == src["missing"].tolist()
    manifest = json.loads((REPO / "data" / "data_manifest.json").read_text())
    sources = table(wb, "Data_Quality", "SourceFiles").set_index("file")
    for name, info in manifest["files"].items():
        assert sources.loc[name, "sha256"] == info["sha256"]
        assert sources.loc[name, "bytes"] == info["bytes"]
        assert sources.loc[name, "source_url"] == info["source_url"]


def test_metric_dictionary_has_every_required_column(wb) -> None:
    md = table(wb, "Metric_Dictionary", "MetricDictionary")
    for column in ("definition", "numerator", "denominator", "exclusions", "unit", "source"):
        assert md[column].notna().all() and (md[column].astype(str).str.len() > 3).all()
    assert {"undiagnosed_share", "auc", "specificity_cost", "holm_adjusted_p"} <= set(md["id"])


def test_five_decision_metrics_are_on_the_executive_sheet(wb) -> None:
    ex = table(wb, "Executive_Summary", "DecisionMetrics")
    assert len(ex) == 5
    est = csv("undiagnosis_estimates").set_index("group")
    assert ex.loc[0, "value"] == pytest.approx(est.loc["All groups", "estimate"], abs=1e-12)
    assert ex["definition"].str.len().min() > 20
    assert "Pipeline refresh" in str(wb["Executive_Summary"]["A4"].value)


def test_every_formula_recalculates_and_every_check_passes(built, wb) -> None:
    values = formula_values(built[1])
    assert len(values) > 200
    assert not [k for k, v in values.items() if isinstance(v, str) and v.startswith("#")]
    assert not [k for k, v in values.items() if v in ("FAIL", "REVIEW")]
    statuses = [v for v in values.values() if v in ("PASS", "ALL CHECKS PASS")]
    assert len(statuses) > 100


def test_calculated_results_are_stored_next_to_the_formulas(built) -> None:
    cached = load_workbook(built[1], data_only=True)
    ws = cached["Executive_Summary"]
    assert table(cached, "Executive_Summary", "DecisionMetrics")["check"].tolist() == ["PASS"] * 5
    status = [ws.cell(r, 2).value for r in range(13, 19)]
    assert status == ["ALL CHECKS PASS"] * 6


def test_committed_workbook_matches_the_committed_tables() -> None:
    config = load_config(root=REPO)
    assert (REPO / WORKBOOK).exists()
    assert verify_workbook(config) == []


def test_a_stale_value_is_detected(tmp_path, built) -> None:
    config, path = built
    stale = tmp_path / "stale.xlsx"
    wb = load_workbook(path)
    ws = wb["Care_Gaps"]
    header = ws.tables["CareGaps"].ref.split(":")[0]
    ws.cell(int(header[1:]) + 1, 6).value = 0.5  # the first group's estimate
    wb.save(stale)
    problems = verify_workbook(config, stale)
    assert problems and "estimate" in problems[0]


def test_a_tampered_estimate_fails_the_workbook_reconciliation(tmp_path) -> None:
    root = tmp_path / "proj"
    (root / "outputs").mkdir(parents=True)
    shutil.copytree(TABLES, root / "outputs" / "tables")
    shutil.copy(REPO / "outputs" / "run_manifest.json", root / "outputs" / "run_manifest.json")
    shutil.copytree(REPO / "configs", root / "configs")
    (root / "data").mkdir()
    shutil.copy(REPO / "data" / "data_manifest.json", root / "data" / "data_manifest.json")
    path = root / "outputs" / "tables" / "undiagnosis_estimates.csv"
    est = pd.read_csv(path)
    est.loc[est["group"] == "All groups", "estimate"] += 0.05
    est.to_csv(path, index=False)
    config = load_config(root=root)
    out = build_workbook(config, root / "tampered.xlsx")
    values = formula_values(out)
    assert "FAIL" in values.values()
    assert "REVIEW" in values.values()
