"""Build excel/nhanes_diabetes_quality_review.xlsx from outputs/tables.

Every displayed value is read from the pipeline's CSV and JSON outputs. Rates, intervals, variances and
reconciliation checks that can be derived from those values are live Excel formulas, so a reviewer can see
them and re-run them. No respondent-level rows are written: every table is a group, model or check summary.
"""

from __future__ import annotations

import json
import re
import shutil
import tempfile
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from nhanes_diabetes.config import Config

WORKBOOK = Path("excel") / "nhanes_diabetes_quality_review.xlsx"
SHEETS = ["Executive_Summary", "Cohort_Flow", "Care_Gaps", "Subgroup_Review", "Model_Tradeoffs",
          "Data_Quality", "Metric_Dictionary"]
ALL = "All groups"
PRIMARY, RACE = "demographic_logistic", "demographic_race_logistic"
ALPHA = 0.05  # display convention for the exploratory labels, not a pipeline result
TOL = 0.000001

HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BANNER_FONT = Font(bold=True, color="9C0006")
WARN_FILL = PatternFill("solid", start_color="FFE599")
BAD_FILL = PatternFill("solid", start_color="F4B6B6")
GOOD_FILL = PatternFill("solid", start_color="C6E0B4")
NOTE_FONT = Font(italic=True, color="555555")
HEADER_ROW = 5

PCT, SPCT, INT, DEC3, DEC2, PVAL = "0.0%", "+0.0%;-0.0%;0.0%", "#,##0", "0.000", "0.00", "0.000"
FORMATS: dict[str, str] = {
    "estimate": PCT, "ci_low": PCT, "ci_high": PCT, "unweighted_share": PCT, "weighted_share_check": PCT,
    "ci_width": PCT, "missing_share": PCT, "missing_share_check": PCT, "share_of_start": PCT,
    "sensitivity": PCT, "sensitivity_ci_low": PCT, "sensitivity_ci_high": PCT, "specificity": PCT,
    "specificity_ci_low": PCT, "specificity_ci_high": PCT, "ppv": PCT, "npv": PCT,
    "false_positive_rate": PCT, "false_positive_rate_ci_low": PCT, "false_positive_rate_ci_high": PCT,
    "specificity_diagnosed": PCT, "specificity_hba1c": PCT, "fpr_diagnosed_label": PCT, "fpr_hba1c_label": PCT,
    "difference": SPCT, "diff_ci_low": SPCT, "diff_ci_high": SPCT, "cost": SPCT, "cost_ci_low": SPCT,
    "cost_ci_high": SPCT, "extra_false_positives_per_100": "+0.0;-0.0;0.0",
    "ratio": DEC2, "ratio_ci_low": DEC2, "ratio_ci_high": DEC2,
    "p_value": PVAL, "p_value_holm": PVAL, "p_value_bootstrap": PVAL,
    "auc_diagnosed_label": DEC3, "auc_diagnosed_ci_low": DEC3, "auc_diagnosed_ci_high": DEC3,
    "auc_hba1c_label": DEC3, "auc_hba1c_ci_low": DEC3, "auc_hba1c_ci_high": DEC3,
    "auc_difference": "+0.000;-0.000;0.000", "auc_diff_ci_low": "+0.000;-0.000;0.000",
    "auc_diff_ci_high": "+0.000;-0.000;0.000",
    "effective_n_kish": "0.0", "effective_n_positives": "0.0", "python_estimate": "0.000000",
    "r_estimate": "0.000000", "effective_n": "0.0",
    "absolute_difference": "0.0E+00", "tolerance": "0.0E+00", "diff_check": "0.0E+00",
    "weighted_denominator": INT, "weighted_undiagnosed": INT, "weighted_population": INT, "bytes": INT,
    "remaining": INT, "removed": INT, "removed_check": INT, "respondents": INT, "hba1c_positive_respondents": INT,
    "undiagnosed_respondents": INT, "respondents_per_group_summary": INT, "diagnosed_among_positive": INT,
    "undiagnosed_among_positive": INT, "positives": INT, "negatives": INT, "missing": INT, "step": "0",
}

Template = str | Callable[[int, int, dict[str, str]], str]


def _cell_value(value: Any) -> Any:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def _banner(ws: Worksheet, banner: str, title: str, note: str) -> None:
    ws["A1"], ws["A1"].font = banner, BANNER_FONT
    ws["A2"], ws["A2"].font = title, Font(bold=True, size=14)
    ws["A3"], ws["A3"].font = note, NOTE_FONT
    ws["A3"].alignment = Alignment(wrap_text=False)


def _write_table(ws: Worksheet, start: int, frame: pd.DataFrame, name: str,
                 formulas: dict[str, Template] | None = None, title: str | None = None,
                 widths: dict[str, float] | None = None) -> dict[str, Any]:
    """Write a titled Excel table whose header sits at ``start``. Returns its geometry for cross references."""
    formulas = formulas or {}
    if title:
        ws.cell(start - 1, 1, title).font = Font(bold=True, size=12)
    columns = list(frame.columns) + list(formulas)
    letters = {c: get_column_letter(i) for i, c in enumerate(columns, start=1)}
    for i, column in enumerate(columns, start=1):
        cell = ws.cell(start, i, column)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        current = ws.column_dimensions[get_column_letter(i)].width or 0
        wanted = (widths or {}).get(column, max(11, min(30, len(column) + 3)))
        ws.column_dimensions[get_column_letter(i)].width = max(current, wanted)
    first = start + 1
    for offset, row in enumerate(frame.itertuples(index=False)):
        r = first + offset
        for i, value in enumerate(row, start=1):
            cell = ws.cell(r, i, _cell_value(value))
            fmt = FORMATS.get(columns[i - 1])
            if fmt:
                cell.number_format = fmt
    for j, (label, template) in enumerate(formulas.items(), start=len(frame.columns) + 1):
        for offset in range(len(frame)):
            r = first + offset
            text = template(r, first, letters) if callable(template) else template.format(r=r, first=first, **letters)
            cell = ws.cell(r, j, text)
            fmt = FORMATS.get(label)
            if fmt:
                cell.number_format = fmt
    last = first + max(len(frame), 1) - 1
    table = Table(displayName=name, ref=f"A{start}:{get_column_letter(len(columns))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    return {"header": start, "first": first, "last": last, "letters": letters, "rows": len(frame)}


def _recon_table(ws: Worksheet, start: int, name: str, title: str,
                 checks: list[tuple[str, str, Any]]) -> dict[str, Any]:
    """Rows of (label, workbook formula, expected value or formula) with difference and PASS/FAIL formulas."""
    header = ["check", "workbook_value", "pipeline_value", "abs_difference", "status"]
    frame = pd.DataFrame({"check": [c[0] for c in checks]})
    geom = _write_table(ws, start, frame, name, title=title, widths={"check": 58})
    for i in range(1, len(header)):
        ws.cell(start, i + 1, header[i]).fill = HEADER_FILL
        ws.cell(start, i + 1).font = HEADER_FONT
    for offset, (_, formula, expected) in enumerate(checks):
        r = geom["first"] + offset
        ws.cell(r, 2, formula).number_format = "#,##0.######"
        ws.cell(r, 3, expected).number_format = "#,##0.######"
        ws.cell(r, 4, f"=ABS(B{r}-C{r})").number_format = "0.0E+00"
        ws.cell(r, 5, f'=IF(D{r}<={TOL},"PASS","FAIL")')
    table = ws.tables[name]
    table.ref = f"A{start}:E{geom['last']}"
    for col in "BCDE":
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 16)
    status = f"E{geom['first']}:E{geom['last']}"
    geom["status_range"] = status
    geom["overall_cell"] = f"B{geom['last'] + 1}"
    ws.cell(geom["last"] + 1, 1, "Overall").font = Font(bold=True)
    ws[geom["overall_cell"]] = f'=IF(COUNTIF({status},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'
    ws[geom["overall_cell"]].font = Font(bold=True)
    return geom


def _flags(ws: Worksheet) -> None:
    rng = "A1:R400"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=BAD_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"REVIEW"'], fill=BAD_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=GOOD_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ALL CHECKS PASS"'], fill=GOOD_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SMALL SAMPLE"'], fill=WARN_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(A1,11)="Exploratory"'], fill=WARN_FILL))


def _page(ws: Worksheet, freeze: str | None = None) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    _flags(ws)


def formula_values(path: Path) -> dict[tuple[str, str], Any]:
    """Recalculate every formula with the independent ``formulas`` engine. Keys are (sheet name, cell)."""
    import formulas

    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    wb = load_workbook(path)
    prefix = f"'[{path.name}]"
    values: dict[tuple[str, str], Any] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    result = solution[f"{prefix}{ws.title.upper()}'!{cell.coordinate}"].value[0, 0]
                    values[(ws.title, cell.coordinate)] = result.item() if hasattr(result, "item") else result
    return values


_FORMULA_CELL = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?)><f>(.*?)</f><v ?/></c>', re.S)


def embed_cached_values(path: Path) -> int:
    """Store each formula's calculated result next to the formula, so previewers that do not calculate show it.

    Excel and LibreOffice still recalculate on open. Returns the number of formulas written.
    """
    values = formula_values(path)
    count = 0
    with tempfile.TemporaryDirectory() as tmp:
        rewritten = Path(tmp) / path.name
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
            names = load_workbook(path).sheetnames
            for item in source.infolist():
                data = source.read(item.filename)
                match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
                if match:
                    sheet = names[int(match.group(1)) - 1]

                    def fill(m: re.Match[str], sheet: str = sheet) -> str:
                        nonlocal count
                        result = values[(sheet, m.group(1))]
                        count += 1
                        head = f'<c r="{m.group(1)}"{m.group(2)}'
                        if isinstance(result, bool):
                            return f'{head} t="b"><f>{m.group(3)}</f><v>{int(result)}</v></c>'
                        if isinstance(result, (int, float)):
                            return f"{head}><f>{m.group(3)}</f><v>{result!r}</v></c>"
                        return f'{head} t="str"><f>{m.group(3)}</f><v>{escape(str(result))}</v></c>'

                    data = _FORMULA_CELL.sub(fill, data.decode("utf-8")).encode("utf-8")
                target.writestr(item, data)
        shutil.copyfile(rewritten, path)
    return count


def load_inputs(config: Config) -> dict[str, Any]:
    tables = {p.stem: pd.read_csv(p) for p in sorted(config.tables_dir.glob("*.csv"))}
    run = json.loads((config.root / "outputs" / "run_manifest.json").read_text(encoding="utf-8"))
    manifest = json.loads(config.manifest_path.read_text(encoding="utf-8"))
    dictionary = yaml.safe_load((config.root / "configs" / "metric_dictionary.yml").read_text(encoding="utf-8"))
    return {"tables": tables, "run": run, "manifest": manifest, "dictionary": dictionary}


def _flag_text(value: Any) -> str:
    return "SMALL SAMPLE" if bool(value) else "OK"


def _status_text(value: Any) -> str:
    return "PASS" if bool(value) else "FAIL"


def care_gap_frame(tables: dict[str, pd.DataFrame], groups: list[str]) -> pd.DataFrame:
    est = tables["undiagnosis_estimates"].set_index("group")
    summary = tables["group_summary"].set_index("group_name")
    rows = []
    for g in groups + [ALL]:
        e, s = est.loc[g], summary.loc[g]
        rows.append({
            "group": g,
            "hba1c_positive_respondents": int(e["respondents"]),
            "undiagnosed_respondents": int(s["undiagnosed_among_positive"]),
            "weighted_denominator": float(e["weighted_denominator"]),
            "weighted_undiagnosed": float(s["weighted_undiagnosed"]),
            "estimate": float(e["estimate"]),
            "ci_low": float(e["ci_low_logit"]),
            "ci_high": float(e["ci_high_logit"]),
            "effective_n_kish": float(e["effective_n_kish"]),
            "small_sample_warning": _flag_text(e["small_n_warning"]),
            "respondents_per_group_summary": int(s["hba1c_positive_respondents"]),
        })
    return pd.DataFrame(rows)


def cohort_flow_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    att = tables["cohort_attrition"].rename(columns={"description": "criterion", "n": "remaining"})
    return att[["criterion", "step", "remaining", "removed"]]


def subgroup_review_frame(tables: dict[str, pd.DataFrame], config: Config) -> pd.DataFrame:
    est = tables["undiagnosis_estimates"].set_index("group")
    con = tables["undiagnosis_contrasts"].set_index("group")
    reference = str(config.analysis("reference_group"))
    rows = []
    for g in config.groups:
        e = est.loc[g]
        row: dict[str, Any] = {
            "group": g, "hba1c_positive_respondents": int(e["respondents"]),
            "effective_n_kish": float(e["effective_n_kish"]), "small_sample_warning": _flag_text(e["small_n_warning"]),
            "estimate": float(e["estimate"]), "ci_low": float(e["ci_low_logit"]), "ci_high": float(e["ci_high_logit"]),
            "difference": None, "diff_ci_low": None, "diff_ci_high": None, "ratio": None, "ratio_ci_low": None,
            "ratio_ci_high": None, "p_value": None, "p_value_holm": None, "analysis_type": None,
        }
        if g != reference:
            c = con.loc[g]
            row.update({"difference": float(c["difference"]), "diff_ci_low": float(c["ci_low"]),
                        "diff_ci_high": float(c["ci_high"]), "ratio": float(c["ratio"]),
                        "ratio_ci_low": float(c["ratio_ci_low"]), "ratio_ci_high": float(c["ratio_ci_high"]),
                        "p_value": float(c["p_value"]), "p_value_holm": float(c["p_value_holm"]),
                        "analysis_type": str(c["analysis_type"])})
        rows.append(row)
    frame = pd.DataFrame(rows)
    order = [reference] + [g for g in config.groups if g != reference]
    return frame.set_index("group").loc[order].reset_index()


def model_comparison_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    spec = tables["model_specification"].set_index("model")
    auc = tables["auc"]
    diff = tables["auc_difference"].set_index("model")
    cost = tables["specificity_cost"]
    rows = []
    for model in spec.index:
        a = auc[(auc["model"] == model) & (auc["evaluated_against"] == "hba1c_criterion")].set_index("label")
        c = cost[(cost["model"] == model) & (cost["group"] == ALL)].iloc[0]
        d = diff.loc[model]
        rows.append({
            "model": model, "role": spec.loc[model, "role"], "description": spec.loc[model, "description"],
            "auc_diagnosed_label": float(a.loc["diagnosed", "auc"]),
            "auc_diagnosed_ci_low": float(a.loc["diagnosed", "ci_low"]),
            "auc_diagnosed_ci_high": float(a.loc["diagnosed", "ci_high"]),
            "auc_hba1c_label": float(a.loc["hba1c_pos", "auc"]),
            "auc_hba1c_ci_low": float(a.loc["hba1c_pos", "ci_low"]),
            "auc_hba1c_ci_high": float(a.loc["hba1c_pos", "ci_high"]),
            "auc_difference": float(d["difference"]), "auc_diff_ci_low": float(d["ci_low"]),
            "auc_diff_ci_high": float(d["ci_high"]), "p_value_bootstrap": float(d["p_value_bootstrap"]),
            "cost": float(c["cost"]), "cost_ci_low": float(c["ci_low"]), "cost_ci_high": float(c["ci_high"]),
        })
    return pd.DataFrame(rows)


def subgroup_performance_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    sub = tables["subgroup_metrics"]
    keep = ["model", "group", "label", "respondents", "positives", "negatives", "effective_n_positives",
            "small_n_warning", "sensitivity", "sensitivity_ci_low", "sensitivity_ci_high", "specificity",
            "specificity_ci_low", "specificity_ci_high", "ppv", "npv"]
    frame = sub[keep].copy()
    frame["small_n_warning"] = frame["small_n_warning"].map(_flag_text)
    return frame.rename(columns={"small_n_warning": "small_sample_warning"}).reset_index(drop=True)


def burden_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    cost = tables["specificity_cost"]
    frame = cost.rename(columns={"cost": "cost", "ci_low": "cost_ci_low", "ci_high": "cost_ci_high"})
    return frame[["model", "group", "specificity_diagnosed", "specificity_hba1c", "cost", "cost_ci_low",
                  "cost_ci_high"]].reset_index(drop=True)


def cohort_by_group_frame(tables: dict[str, pd.DataFrame], groups: list[str]) -> pd.DataFrame:
    summary = tables["group_summary"].set_index("group_name")
    rows = []
    for g in groups + [ALL]:
        s = summary.loc[g]
        rows.append({"group": g, "respondents": int(s["respondents"]),
                     "weighted_population": float(s["weighted_population"]),
                     "hba1c_positive_respondents": int(s["hba1c_positive_respondents"]),
                     "diagnosed_among_positive": int(s["diagnosed_among_positive"]),
                     "undiagnosed_among_positive": int(s["undiagnosed_among_positive"])})
    return pd.DataFrame(rows)


def missingness_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    return tables["missingness_by_group"][["group", "variable", "respondents", "missing", "missing_share"]].copy()


def checks_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    q = tables["data_quality"].copy()
    q["passed"] = q["passed"].map(_status_text)
    return q.rename(columns={"passed": "status", "check_name": "check"})[["suite", "check", "status", "detail"]]


def sources_frame(manifest: dict[str, Any]) -> pd.DataFrame:
    rows = [{"file": name, "source_url": info["source_url"], "bytes": int(info["bytes"]),
             "sha256": info["sha256"], "retrieved_at": info["retrieved_at"],
             "columns_verified": ", ".join(info["columns_verified"])}
            for name, info in sorted(manifest["files"].items())]
    return pd.DataFrame(rows)


def crosscheck_frame(tables: dict[str, pd.DataFrame]) -> pd.DataFrame | None:
    x = tables.get("r_python_crosscheck")
    if x is None or x.empty:
        return None
    frame = x[["quantity", "group", "python", "r", "absolute_difference", "tolerance"]].copy()
    return frame.rename(columns={"python": "python_estimate", "r": "r_estimate"})


def dictionary_frame(dictionary: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{k: m[k] for k in ("id", "name", "definition", "numerator", "denominator", "exclusions",
                                            "unit", "source")} for m in dictionary["metrics"]])


def _ref(sheet: str, geom: dict[str, Any], column: str, row: int) -> str:
    return f"{sheet}!{geom['letters'][column]}{row}"


def build_workbook(config: Config, output: Path | None = None) -> Path:
    inputs = load_inputs(config)
    tables, run, manifest, dictionary = inputs["tables"], inputs["run"], inputs["manifest"], inputs["dictionary"]
    groups = config.groups
    banner = " ".join(str(dictionary["banner"]).split())
    refresh = str(run["generated_at"])
    retrieved = sorted({info["retrieved_at"] for info in manifest["files"].values()})
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    sheets = {name: wb.create_sheet(name) for name in SHEETS}
    created = datetime.fromisoformat(refresh)
    wb.properties.created = wb.properties.modified = created.replace(tzinfo=None)
    wb.properties.creator = wb.properties.lastModifiedBy = "nhanes-diabetes"
    wb.properties.title = "NHANES 2017-2018 diabetes care-gap quality review"
    wb.calculation.fullCalcOnLoad = True

    # Care_Gaps
    ws = sheets["Care_Gaps"]
    _banner(ws, banner, "Care gaps: undiagnosed diabetes among people meeting the HbA1c criterion",
            "Weighted estimate with a 95% logit interval, unweighted n and Kish effective n. Rows flagged SMALL SAMPLE have wide intervals.")
    cg = care_gap_frame(tables, groups)
    care = _write_table(ws, HEADER_ROW, cg, "CareGaps", widths={"group": 24, "small_sample_warning": 20}, formulas={
        "unweighted_share": "={undiagnosed_respondents}{r}/{hba1c_positive_respondents}{r}",
        "weighted_share_check": "={weighted_undiagnosed}{r}/{weighted_denominator}{r}",
        "ci_width": "={ci_high}{r}-{ci_low}{r}",
        "estimate_reconciles": '=IF(ABS({weighted_share_check}{r}-{estimate}{r})<=' + str(TOL) + ',"PASS","FAIL")',
        "n_reconciles": '=IF({hba1c_positive_respondents}{r}={respondents_per_group_summary}{r},"PASS","FAIL")',
    })
    care_row = {g: care["first"] + i for i, g in enumerate(groups + [ALL])}
    ws.cell(care["last"] + 2, 1, "Weighted undiagnosed share and people counts reconcile to group_summary.csv and "
            "undiagnosis_estimates.csv. Unweighted share is shown for context only.").font = NOTE_FONT
    _page(ws, f"B{care['first']}")

    # Cohort_Flow
    ws = sheets["Cohort_Flow"]
    _banner(ws, banner, "Cohort flow: from all NHANES participants to the analysis cohort",
            "Each step's remaining count comes from cohort_attrition.csv. Removed is re-derived by formula and reconciled below.")
    att = cohort_flow_frame(tables)
    flow = _write_table(ws, HEADER_ROW, att, "CohortFlow", widths={"criterion": 62}, formulas={
        "removed_check": lambda r, first, L: "=0" if r == first else f"={L['remaining']}{r - 1}-{L['remaining']}{r}",
        "share_of_start": "={remaining}{r}/{remaining}$" + str(HEADER_ROW + 1),
        "removed_reconciles": '=IF({removed}{r}={removed_check}{r},"PASS","FAIL")',
    })
    cbg = cohort_by_group_frame(tables, groups)
    grp_start = flow["last"] + 5
    grp = _write_table(ws, grp_start, cbg, "CohortByGroup", title="Analysis cohort by group (summary counts, no respondent rows)")
    g_first, g_last = grp["first"], grp["last"] - 1  # last row is All groups
    all_row = grp["last"]
    L = grp["letters"]
    counts = run["sample_counts"]
    fl = flow["letters"]
    recon_start = grp["last"] + 5
    recon = _recon_table(ws, recon_start, "CohortReconciliation", "Reconciliation", [
        ("First step equals participants in run_manifest.json", f"={fl['remaining']}{flow['first']}", int(counts["participants"])),
        ("Last step equals analysis cohort in run_manifest.json", f"={fl['remaining']}{flow['last']}", int(counts["analysis_cohort"])),
        ("Sum of group respondents equals the last step", f"=SUM({L['respondents']}{g_first}:{L['respondents']}{g_last})",
         f"={fl['remaining']}{flow['last']}"),
        ("All groups row equals the sum of groups", f"={L['respondents']}{all_row}",
         f"=SUM({L['respondents']}{g_first}:{L['respondents']}{g_last})"),
        ("Diagnosed plus undiagnosed equals HbA1c positives (All groups)",
         f"={L['diagnosed_among_positive']}{all_row}+{L['undiagnosed_among_positive']}{all_row}",
         f"={L['hba1c_positive_respondents']}{all_row}"),
        ("Total removed equals first minus last step", f"=SUM({fl['removed']}{flow['first']}:{fl['removed']}{flow['last']})",
         f"={fl['remaining']}{flow['first']}-{fl['remaining']}{flow['last']}"),
    ])
    _page(ws, f"A{flow['first']}")

    # Subgroup_Review
    ws = sheets["Subgroup_Review"]
    _banner(ws, banner, "Subgroup review: undiagnosed share by group and comparison with the reference group",
            f"Only the pre-specified comparison is confirmatory. Every other comparison is exploratory and labelled with its Holm-adjusted p-value (labels use {ALPHA} as a display convention).")
    sr = subgroup_review_frame(tables, config)
    review = _write_table(ws, HEADER_ROW, sr, "SubgroupReview", widths={"group": 24, "small_sample_warning": 20}, formulas={
        "review_label": ('=IF({analysis_type}{r}="","Reference group",IF({analysis_type}{r}="confirmatory",'
                         '"Pre-specified comparison","Exploratory: "&IF({p_value_holm}{r}<' + str(ALPHA) +
                         ',"below ' + str(ALPHA) + ' after Holm adjustment","not below ' + str(ALPHA) + ' after Holm adjustment")))'),
        "interval_excludes_zero": '=IF({difference}{r}="","",IF(OR({diff_ci_low}{r}>0,{diff_ci_high}{r}<0),"YES","NO"))',
        "holm_at_least_raw_p": '=IF({p_value}{r}="","",IF({p_value_holm}{r}>={p_value}{r}-' + str(TOL) + ',"PASS","FAIL"))',
    })
    ws.column_dimensions[review["letters"]["review_label"]].width = 52
    review_row = {g: review["first"] + i for i, g in enumerate(sr["group"])}
    ws.cell(review["last"] + 2, 1, "Difference is the group's undiagnosed share minus the reference group's, in percentage points. "
            "Intervals are 95%.").font = NOTE_FONT
    _page(ws, f"B{review['first']}")

    # Model_Tradeoffs
    ws = sheets["Model_Tradeoffs"]
    _banner(ws, banner, "Model tradeoffs: does the training label change who a screening model misses or over-flags?",
            "Both labels are scored against the HbA1c criterion. False-positive rate is one minus specificity and is a formula. Not clinical risk models.")
    mc = model_comparison_frame(tables)
    comparison = _write_table(ws, HEADER_ROW, mc, "ModelComparison", widths={"model": 26, "role": 24, "description": 46}, formulas={
        "auc_difference_reconciles": '=IF(ABS(({auc_diagnosed_label}{r}-{auc_hba1c_label}{r})-{auc_difference}{r})<=' + str(TOL) + ',"PASS","FAIL")',
    })
    perf_start = comparison["last"] + 5
    sp = subgroup_performance_frame(tables)
    perf = _write_table(ws, perf_start, sp, "SubgroupPerformance", title="Sensitivity, specificity and false-positive rate by model, group and label",
                        widths={"group": 24}, formulas={
                            "false_positive_rate": "=1-{specificity}{r}",
                            "false_positive_rate_ci_low": "=1-{specificity_ci_high}{r}",
                            "false_positive_rate_ci_high": "=1-{specificity_ci_low}{r}",
                        })
    burden_start = perf["last"] + 5
    bf = burden_frame(tables)
    burden = _write_table(ws, burden_start, bf, "FalsePositiveBurden",
                          title="False-positive burden of the HbA1c label versus the diagnosed label",
                          widths={"group": 24}, formulas={
                              "fpr_diagnosed_label": "=1-{specificity_diagnosed}{r}",
                              "fpr_hba1c_label": "=1-{specificity_hba1c}{r}",
                              "extra_false_positives_per_100": "=({fpr_hba1c_label}{r}-{fpr_diagnosed_label}{r})*100",
                              "cost_reconciles": '=IF(ABS(({specificity_diagnosed}{r}-{specificity_hba1c}{r})-{cost}{r})<=' + str(TOL) + ',"PASS","FAIL")',
                          })
    ws.cell(burden["last"] + 2, 1, "Extra false positives per 100 is the change in false positives per 100 HbA1c-negative people when the model "
            "is trained on the HbA1c label instead of the diagnosed label. Positive means more false positives.").font = NOTE_FONT
    burden_row = {(m, g): burden["first"] + i for i, (m, g) in enumerate(zip(bf["model"], bf["group"], strict=True))}
    _page(ws, f"D{comparison['first']}")

    # Data_Quality
    ws = sheets["Data_Quality"]
    _banner(ws, banner, "Data quality: pipeline checks, missingness, duplicates and source file hashes",
            "Checks come from data_quality.csv (SQL suites run on every build). Source hashes come from data/data_manifest.json.")
    checks = _write_table(ws, HEADER_ROW, checks_frame(tables), "QualityChecks", widths={"suite": 18, "check": 36, "detail": 60})
    ws.cell(checks["last"] + 1, 1, "Checks passing").font = Font(bold=True)
    passing_cell = f"C{checks['last'] + 1}"
    ws[passing_cell] = f'=COUNTIF({checks["letters"]["status"]}{checks["first"]}:{checks["letters"]["status"]}{checks["last"]},"PASS")&" of "&ROWS({checks["letters"]["status"]}{checks["first"]}:{checks["letters"]["status"]}{checks["last"]})'
    miss_start = checks["last"] + 5
    miss = _write_table(ws, miss_start, missingness_frame(tables), "Missingness",
                        title="Missingness by group and input variable (refused, do not know, not asked or blank)",
                        formulas={
                            "missing_share_check": "={missing}{r}/{respondents}{r}",
                            "missing_reconciles": '=IF(ABS({missing_share_check}{r}-{missing_share}{r})<=' + str(TOL) + ',"PASS","FAIL")',
                        })
    src_start = miss["last"] + 5
    src = _write_table(ws, src_start, sources_frame(manifest), "SourceFiles",
                       title="Source files: CDC NHANES 2017-2018 public-use data (hashes verified on download)",
                       widths={"file": 12, "source_url": 60, "sha256": 66, "columns_verified": 46})
    cross = crosscheck_frame(tables)
    if cross is not None:
        cross_start = src["last"] + 5
        _write_table(ws, cross_start, cross, "RCrosscheck",
                                  title="Independent check: Python survey estimates against R's survey package",
                                  formulas={
                                      "diff_check": "=ABS({python_estimate}{r}-{r_estimate}{r})",
                                      "within_tolerance": '=IF({diff_check}{r}<={tolerance}{r},"PASS","FAIL")',
                                  })
    _page(ws, f"A{checks['first']}")

    # Metric_Dictionary
    ws = sheets["Metric_Dictionary"]
    _banner(ws, banner, "Metric dictionary: numerator, denominator, exclusions, unit and source",
            "Definitions live in configs/metric_dictionary.yml and are shared with the stakeholder brief and the Tableau field dictionary.")
    md = dictionary_frame(dictionary)
    md_geom = _write_table(ws, HEADER_ROW, md, "MetricDictionary", widths={
        "id": 24, "name": 30, "definition": 48, "numerator": 48, "denominator": 40, "exclusions": 44, "unit": 40, "source": 46})
    for row in ws.iter_rows(min_row=md_geom["first"], max_row=md_geom["last"]):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _page(ws, f"C{md_geom['first']}")

    # Executive_Summary (last, because it references the other sheets)
    ws = sheets["Executive_Summary"]
    _banner(ws, banner, "Executive summary: five decision metrics for diabetes care-gap monitoring",
            "Values are read from the pipeline outputs; the Check column re-reads the same value from the detail sheet.")
    est = tables["undiagnosis_estimates"].set_index("group")
    gs = tables["group_summary"].set_index("group_name")
    con = tables["undiagnosis_contrasts"].set_index("group")
    cost = tables["specificity_cost"]
    comparison_group, reference = str(config.analysis("comparison_group")), str(config.analysis("reference_group"))
    race_black = cost[(cost["model"] == RACE) & (cost["group"] == comparison_group)].iloc[0]
    sub = tables["subgroup_metrics"]
    negatives = sub[(sub["model"] == RACE) & (sub["label"] == "diagnosed") & (sub["group"] == comparison_group)].iloc[0]
    defs = {m["id"]: m for m in dictionary["metrics"]}

    def definition(metric_id: str) -> str:
        return str(defs[metric_id]["definition"])

    share_unit = "share of HbA1c-positive people"
    rows = [
        {"metric": "Undiagnosed share, all groups", "value": float(est.loc[ALL, "estimate"]),
         "ci_low": float(est.loc[ALL, "ci_low_logit"]), "ci_high": float(est.loc[ALL, "ci_high_logit"]),
         "unit": share_unit, "n_basis": "HbA1c-positive respondents", "unweighted_n": int(est.loc[ALL, "respondents"]),
         "effective_n": float(est.loc[ALL, "effective_n_kish"]), "definition": definition("undiagnosed_share"),
         "source": "undiagnosis_estimates.csv"},
        {"metric": "People with undiagnosed diabetes, weighted", "value": float(gs.loc[ALL, "weighted_undiagnosed"]),
         "ci_low": None, "ci_high": None, "unit": "people (survey-weighted)", "n_basis": "undiagnosed respondents",
         "unweighted_n": int(gs.loc[ALL, "undiagnosed_among_positive"]), "effective_n": None,
         "definition": definition("weighted_undiagnosed_people"), "source": "group_summary.csv"},
        {"metric": f"Undiagnosed share, {comparison_group}", "value": float(est.loc[comparison_group, "estimate"]),
         "ci_low": float(est.loc[comparison_group, "ci_low_logit"]), "ci_high": float(est.loc[comparison_group, "ci_high_logit"]),
         "unit": share_unit, "n_basis": "HbA1c-positive respondents", "unweighted_n": int(est.loc[comparison_group, "respondents"]),
         "effective_n": float(est.loc[comparison_group, "effective_n_kish"]), "definition": definition("undiagnosed_share"),
         "source": "undiagnosis_estimates.csv"},
        {"metric": f"Difference versus {reference} (pre-specified)", "value": float(con.loc[comparison_group, "difference"]),
         "ci_low": float(con.loc[comparison_group, "ci_low"]), "ci_high": float(con.loc[comparison_group, "ci_high"]),
         "unit": "percentage points", "n_basis": "HbA1c-positive respondents, comparison group",
         "unweighted_n": int(est.loc[comparison_group, "respondents"]),
         "effective_n": float(est.loc[comparison_group, "effective_n_kish"]), "definition": definition("undiagnosed_gap"),
         "source": "undiagnosis_contrasts.csv"},
        {"metric": f"Specificity cost of the HbA1c label, {comparison_group}, race-included model",
         "value": float(race_black["cost"]), "ci_low": float(race_black["ci_low"]), "ci_high": float(race_black["ci_high"]),
         "unit": "percentage points", "n_basis": "HbA1c-negative respondents",
         "unweighted_n": int(negatives["negatives"]), "effective_n": float(negatives["effective_n_negatives"]),
         "definition": definition("specificity_cost"), "source": "specificity_cost.csv"},
    ]
    dm = pd.DataFrame(rows)
    ws["A4"], ws["A4"].font = f"Pipeline refresh: {refresh}. Source files retrieved: {', '.join(retrieved)}. Code commit: {run['code_commit']}.", NOTE_FONT
    burden_ref = burden_row[(RACE, comparison_group)]
    summary = _write_table(ws, HEADER_ROW, dm, "DecisionMetrics", widths={"metric": 54, "unit": 28, "n_basis": 30, "definition": 70, "source": 26}, formulas={
        "check": lambda r, first, Lt: [
            f'=IF(ABS({Lt["value"]}{r}-{_ref("Care_Gaps", care, "estimate", care_row[ALL])})<={TOL},"PASS","FAIL")',
            f'=IF(ABS({Lt["value"]}{r}-{_ref("Care_Gaps", care, "weighted_undiagnosed", care_row[ALL])})<=0.5,"PASS","FAIL")',
            f'=IF(ABS({Lt["value"]}{r}-{_ref("Care_Gaps", care, "estimate", care_row[comparison_group])})<={TOL},"PASS","FAIL")',
            f'=IF(ABS({Lt["value"]}{r}-{_ref("Subgroup_Review", review, "difference", review_row[comparison_group])})<={TOL},"PASS","FAIL")',
            f'=IF(ABS({Lt["value"]}{r}-{_ref("Model_Tradeoffs", burden, "cost", burden_ref)})<={TOL},"PASS","FAIL")',
        ][r - first],
    })
    for offset, fmt in enumerate([PCT, INT, PCT, SPCT, SPCT]):
        r = summary["first"] + offset
        for col in ("value", "ci_low", "ci_high"):
            ws[f"{summary['letters'][col]}{r}"].number_format = fmt
    def_col = list(dm.columns).index("definition") + 1
    for row in ws.iter_rows(min_row=summary["first"], max_row=summary["last"], min_col=def_col, max_col=def_col):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    status_start = summary["last"] + 3
    ws.cell(status_start - 1, 1, "Reconciliation status by sheet").font = Font(bold=True, size=12)
    status_rows = [
        ("Cohort_Flow", f"=Cohort_Flow!{recon['overall_cell']}"),
        ("Care_Gaps", f'=IF(COUNTIF(Care_Gaps!{care["letters"]["estimate_reconciles"]}{care["first"]}:{care["letters"]["n_reconciles"]}{care["last"]},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'),
        ("Subgroup_Review", f'=IF(COUNTIF(Subgroup_Review!{review["letters"]["holm_at_least_raw_p"]}{review["first"]}:{review["letters"]["holm_at_least_raw_p"]}{review["last"]},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'),
        ("Model_Tradeoffs", f'=IF(COUNTIF(Model_Tradeoffs!{burden["letters"]["cost_reconciles"]}{burden["first"]}:{burden["letters"]["cost_reconciles"]}{burden["last"]},"FAIL")+COUNTIF(Model_Tradeoffs!{comparison["letters"]["auc_difference_reconciles"]}{comparison["first"]}:{comparison["letters"]["auc_difference_reconciles"]}{comparison["last"]},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'),
        ("Data_Quality", f'=IF(COUNTIF(Data_Quality!{checks["letters"]["status"]}{checks["first"]}:{checks["letters"]["status"]}{checks["last"]},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'),
        ("Executive_Summary", f'=IF(COUNTIF({summary["letters"]["check"]}{summary["first"]}:{summary["letters"]["check"]}{summary["last"]},"FAIL")=0,"ALL CHECKS PASS","REVIEW")'),
    ]
    for i, (label, formula) in enumerate(status_rows):
        ws.cell(status_start + i, 1, label)
        ws.cell(status_start + i, 2, formula)
    lim_start = status_start + len(status_rows) + 2
    ws.cell(lim_start - 1, 1, "Population").font = Font(bold=True, size=12)
    ws.cell(lim_start, 1, "The analysis cohort represents " + " ".join(str(dictionary["population"]).split()) + ".")
    ws.cell(lim_start + 2, 1, "Limitations").font = Font(bold=True, size=12)
    for i, text in enumerate(dictionary["limitations"]):
        ws.cell(lim_start + 3 + i, 1, f"{i + 1}. {text}")
    _page(ws, f"A{summary['first']}")

    target = output or (config.root / WORKBOOK)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    embed_cached_values(target)
    return target


def expected_frames(config: Config) -> dict[str, tuple[str, pd.DataFrame]]:
    """Sheet name to (first header cell, expected data columns) used to verify a saved workbook."""
    inputs = load_inputs(config)
    tables = inputs["tables"]
    groups = config.groups
    frames: dict[str, tuple[str, pd.DataFrame]] = {
        "Cohort_Flow": ("criterion", cohort_flow_frame(tables)),
        "Care_Gaps": ("group", care_gap_frame(tables, groups)),
        "Subgroup_Review": ("group", subgroup_review_frame(tables, config)),
        "Model_Tradeoffs": ("model", model_comparison_frame(tables)),
        "Data_Quality": ("suite", checks_frame(tables)),
        "Metric_Dictionary": ("id", dictionary_frame(inputs["dictionary"])),
    }
    return frames


def verify_workbook(config: Config, path: Path | None = None) -> list[str]:
    """Compare stored values in a saved workbook with the current pipeline tables. One message per mismatch."""
    path = path or (config.root / WORKBOOK)
    if not path.exists():
        return [f"{path} does not exist"]
    wb = load_workbook(path)
    problems = [f"missing sheet {s}" for s in SHEETS if s not in wb.sheetnames]
    if problems:
        return problems
    for sheet, (key, frame) in expected_frames(config).items():
        ws = wb[sheet]
        header = next((r for r in range(1, 20) if ws.cell(r, 1).value == key), None)
        if header is None:
            problems.append(f"{sheet}: header row not found")
            continue
        columns = {str(ws.cell(header, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(header, c).value}
        for offset, row in enumerate(frame.to_dict("records")):
            for name, value in row.items():
                stored = ws.cell(header + 1 + offset, columns[name]).value
                expected = _cell_value(value)
                if isinstance(expected, float) and isinstance(stored, (int, float)):
                    if abs(stored - expected) > 1e-9 * max(1.0, abs(expected)):
                        problems.append(f"{sheet} row {offset + 1} {name}: {stored} != {expected}")
                elif stored != expected:
                    problems.append(f"{sheet} row {offset + 1} {name}: {stored!r} != {expected!r}")
    return problems
